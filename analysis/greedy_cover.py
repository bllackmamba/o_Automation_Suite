"""
Order-independence test for the R greedy set-cover (read-only analysis).

Best-marginal-gain greedy: at each step, over ALL remaining Repeat_R rows, pick
the row that adds the most NEW coverage of the Repeat pool; break ties by a
configurable order. Run under several tie-break variants to see whether the
survivor count / covering set is stable (a real "wall") or an artifact of one
arbitrary tie-break.

Definitions (locked with Tai, 2026-08-08):
  - Universe  = the 4,882,437 Main Data 6-combos with match_count(combo, D4687)>=1
                (Rule Set 1 Repeat pool; D4687 = {3,6,9,14,21,22}).
  - Candidate = each of the 6,885 Repeat_R rows (a set of 6..27 numbers).
  - Coverage  = CONTAINMENT: R row R covers combo e iff e ⊆ set(R)
                (i.e. match_count(R, e) == 6). Validated cell-exact against
                Match_Matrix_969.xlsx on the 969 real draws before this run.

Every selection step calls the standing Rule-1 checker (analysis.pool_invariant)
on the newly covered combos — a hard failure aborts immediately.

Read-only. `python3 -m analysis.greedy_cover`.
"""
from __future__ import annotations

import heapq
import math
import sys
import time
from itertools import combinations
from pathlib import Path

import numpy as np
import openpyxl

from analysis.pool_invariant import (assert_pool_invariant_array,
                                      REFGROUP_D4687, PoolInvariantViolation)

REPO = Path(__file__).resolve().parent.parent
RXLSX = REPO.parent / "info" / "R_corrected_D4687_v2.xlsx"
POOL, PICK = 45, 6
N = math.comb(POOL, PICK)                       # 8,145,060
D4687 = sorted(REFGROUP_D4687)                  # [3,6,9,14,21,22]
MILESTONES = (0.50, 0.90, 0.95, 0.99, 1.00)
SAMPLE_UNRANK = 2000                            # per-step number-space spot check


# ── combinadic rank / unrank (colex bijection, k=6) ──────────────────────────
def rank6(c) -> int:
    """Rank a sorted 1-indexed 6-combo into [0, C(45,6))."""
    return (math.comb(c[0]-1,1)+math.comb(c[1]-1,2)+math.comb(c[2]-1,3)
           +math.comb(c[3]-1,4)+math.comb(c[4]-1,5)+math.comb(c[5]-1,6))


def unrank6(r: int) -> tuple[int, ...]:
    """Inverse of rank6 → sorted 1-indexed 6-combo."""
    out = []
    for i in range(6, 0, -1):
        # largest v with C(v, i) <= r
        v = i - 1
        while math.comb(v + 1, i) <= r:
            v += 1
        out.append(v + 1)                       # back to 1-indexed
        r -= math.comb(v, i)
    return tuple(sorted(out))


# ── load Repeat_R rows ───────────────────────────────────────────────────────
def load_repeat_r():
    wb = openpyxl.load_workbook(RXLSX, read_only=True, data_only=True)
    ws = wb["Repeat_R"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ix = {h: i for i, h in enumerate(hdr)}
    pos = [i for h, i in ix.items() if isinstance(h, str) and h.startswith("pos_")]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        nums = sorted({int(r[i]) for i in pos if r[i] not in (None, "", "None")})
        wlabel = str(r[ix["Set_Label"]])
        wnum = int(wlabel[1:]) if wlabel[1:].isdigit() else 10**9
        rows.append({"nums": nums, "wnum": wnum,
                     "len": r[ix["Row_Length"]], "label": wlabel})
    wb.close()
    return rows


# ── precompute per-row in-universe coverage ranks ────────────────────────────
def build_coverage(rows):
    """Per-row int32 array of ranks of its 6-subsets that lie in U (share >=1
    with D4687). Also returns the covered-union size (the greedy target)."""
    d4687 = set(D4687)
    cover = []
    union = np.zeros(N, dtype=bool)
    t0 = time.time(); cells = 0
    for k, row in enumerate(rows):
        ranks = []
        for c in combinations(row["nums"], 6):
            if d4687.isdisjoint(c):
                continue                        # subset not in Repeat pool → skip
            ranks.append(rank6(c))
            cells += 1
        a = np.asarray(ranks, dtype=np.int32)
        cover.append(a)
        union[a] = True
        if (k + 1) % 1000 == 0:
            print(f"  coverage: {k+1}/{len(rows)} rows, {cells:,} cells, "
                  f"{time.time()-t0:.0f}s", flush=True)
    return cover, int(union.sum())


# ── tie-break variants → per-row integer key (lower = preferred on ties) ──────
def tiebreak_keys(rows):
    n = len(rows)
    order_numlex = sorted(range(n), key=lambda i: rows[i]["nums"])
    numlex_rank = [0] * n
    for pos_, i in enumerate(order_numlex):
        numlex_rank[i] = pos_
    return {
        "V1_gen_rainbow":      [rows[i]["wnum"] for i in range(n)],
        "V2_gen_rainbow_rev":  [-rows[i]["wnum"] for i in range(n)],
        "V3_len_asc_then_w":   [rows[i]["len"] * 10**7 + rows[i]["wnum"] for i in range(n)],
        "V4_numbers_lex":      [numlex_rank[i] for i in range(n)],
    }


# ── lazy best-marginal-gain greedy under one tie-break ───────────────────────
def greedy(cover, target, tb_key, rows, variant):
    covered = np.zeros(N, dtype=bool)
    heap = [(-cover[i].size, tb_key[i], i) for i in range(len(cover))]
    heapq.heapify(heap)
    picked = []
    milestones = {}
    total = 0
    t0 = time.time()
    while heap and total < target:
        neg_g, tb, i = heapq.heappop(heap)
        row_ranks = cover[i]
        new_mask = ~covered[row_ranks]
        gain = int(new_mask.sum())
        if gain < -neg_g:                        # stale → re-insert with true gain
            heapq.heappush(heap, (-gain, tb, i))
            continue
        if gain == 0:
            break                                # nothing left to add
        # ── SELECT row i ──
        new_ranks = row_ranks[new_mask]
        # standing Rule-1 invariant on the newly covered combos (Repeat stream)
        smp = new_ranks if new_ranks.size <= SAMPLE_UNRANK else \
            new_ranks[np.linspace(0, new_ranks.size - 1, SAMPLE_UNRANK).astype(int)]
        arr = np.array([unrank6(int(r)) for r in smp], dtype=np.int32)
        assert_pool_invariant_array(
            arr, "Repeat",
            step_label=f"{variant} step {len(picked)+1} (row {rows[i]['label']})")
        covered[new_ranks] = True
        total += gain
        picked.append(i)
        frac = total / target
        for m in MILESTONES:
            if m not in milestones and frac >= m:
                milestones[m] = len(picked)
    return {"survivors": len(picked), "picked": picked, "covered": total,
            "target": target, "milestones": milestones, "secs": time.time() - t0}


def main():
    print(f"Loading Repeat_R from {RXLSX.name} …", flush=True)
    rows = load_repeat_r()
    print(f"  {len(rows)} rows, |set| {min(r['len'] for r in rows)}–"
          f"{max(r['len'] for r in rows)}", flush=True)
    print("Building in-universe coverage (containment) …", flush=True)
    cover, target = build_coverage(rows)
    U = math.comb(45, 6) - math.comb(45 - 6, 6)
    print(f"\nUniverse U (Repeat pool, match(D4687)>=1) = {U:,}")
    print(f"Covered-union of all Repeat_R subsets (greedy target) = {target:,}"
          f"  ({target/U:.4%} of U)", flush=True)

    keys = tiebreak_keys(rows)
    results = {}
    for variant, tb in keys.items():
        print(f"\n── greedy variant {variant} ──", flush=True)
        try:
            res = greedy(cover, target, tb, rows, variant)
        except PoolInvariantViolation as e:
            print(f"  !!! INVARIANT VIOLATION: {e}", flush=True)
            results[variant] = {"violation": str(e)}
            continue
        ms = "  ".join(f"{int(m*100)}%={res['milestones'].get(m,'—')}"
                       for m in MILESTONES)
        print(f"  survivors (full cover) = {res['survivors']}  "
              f"covered {res['covered']:,}/{res['target']:,}  {res['secs']:.0f}s")
        print(f"  rows to reach: {ms}", flush=True)
        results[variant] = res

    # ── cross-variant stability ──
    print("\n" + "=" * 70)
    print("CROSS-VARIANT STABILITY")
    good = {v: r for v, r in results.items() if "violation" not in r}
    print("survivor counts:", {v: r["survivors"] for v, r in good.items()})
    vs = list(good)
    print("\npairwise selected-set Jaccard overlap:")
    for a in range(len(vs)):
        for b in range(a + 1, len(vs)):
            A, B = set(good[vs[a]]["picked"]), set(good[vs[b]]["picked"])
            j = len(A & B) / len(A | B) if (A | B) else 1.0
            print(f"  {vs[a]} vs {vs[b]}: |A|={len(A)} |B|={len(B)} "
                  f"∩={len(A&B)} Jaccard={j:.3f}")
    print("\nmilestone rows-to-coverage across variants:")
    for m in MILESTONES:
        print(f"  {int(m*100)}%:",
              {v: r["milestones"].get(m) for v, r in good.items()})


if __name__ == "__main__":
    main()
