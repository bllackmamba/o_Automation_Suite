"""
R classification verification — spread tiers + shallow anchor (read-only).

Ground-truth pass for two LOCKED-rule classifiers that did NOT exist on disk as
runnable code (only a persisted Constituent_Groups column read by
info/regenerate_r_sheet1.py, and a pass-through stub / frozen-D4687-list version
in syndicate_core/escalation.py). Both are REBUILT here from the real primitives
in syndicate_core.stacked_blocks (since_last_map / block_layout) and validated by
hard gates before any tabulation.

Reference choice:
  - shallow anchor + full-history spread tiers → OWN-ERA rolling (each draw vs
    its own predecessor era) — a draw's spread only means something relative to
    its own point in history.
  - the spread COUNTING MECHANISM is separately validated in FIXED-D4687 mode
    against R's persisted True_Spread (GATE C) — a mechanism check only; it is
    expected to diverge once switched to own-era for draws, which is not a bug.

Datasets (found on disk):
  - History: Games/SAT/Variable_inputs_sat/Base_sat/B1_sat_updated.csv (970 draws,
    newest-first w1..w970).
  - R rows : ../info/R_corrected_D4687_v2.xlsx  (persisted True_Spread).

Findings summary in docs/R_SPREAD_SHALLOW_VERIFICATION.md.
Read-only: touches no tracked code/CSV. `python3 -m analysis.r_spread_shallow`.
"""
from __future__ import annotations
import csv, math
from itertools import combinations
from pathlib import Path
from collections import Counter

from syndicate_core.stacked_blocks import since_last_map, block_layout

REPO = Path(__file__).resolve().parent.parent
INFO = REPO.parent / "info"
B1 = REPO / "Games/SAT/Variable_inputs_sat/Base_sat/B1_sat_updated.csv"
RXLSX = INFO / "R_corrected_D4687_v2.xlsx"
POOL, PICK = 45, 6
C_POOL_PICK = math.comb(POOL, PICK)

D4687 = {3, 6, 9, 14, 21, 22}
DOC_SHALLOW_25 = {3,6,8,9,10,11,12,14,16,19,20,21,22,23,25,28,30,31,32,33,36,39,40,43,44}
KNOWN_EXCEPTIONS = {"w115", "w308", "w362", "w514", "w733", "w811"}
N_SHALLOW_BLOCKS = 6   # Rule 4 "recency-ordinal 1-6" (SL0 = ref draw's own nums, included)


# ── data ──────────────────────────────────────────────────────────────────────
def load_history():
    """Newest-first list of {'w','draw','nums'} from B1_sat_updated.csv."""
    rows = []
    with B1.open() as f:
        for r in csv.DictReader(f):
            nums = {int(r[f"pos_{i}"]) for i in range(1, PICK + 1) if r.get(f"pos_{i}")}
            rows.append({"w": r["w"].strip(),
                         "draw": (r.get("update") or "").strip(), "nums": nums})
    return rows


def load_r_spreads(sheet):
    """Persisted True_Spread values from one sheet of R_corrected_D4687_v2.xlsx."""
    import openpyxl
    wb = openpyxl.load_workbook(RXLSX, read_only=True, data_only=True)
    ws = wb[sheet]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ix = {h: i for i, h in enumerate(hdr)}
    pos_cols = [i for h, i in ix.items() if isinstance(h, str) and h.startswith("pos_")]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        nums = set()
        for i in pos_cols:
            v = r[i]
            if v not in (None, "", "None"):
                try: nums.add(int(float(v)))
                except ValueError: pass
        out.append({"nums": nums, "True_Spread": r[ix["True_Spread"]]})
    wb.close()
    return out


# ── classifiers (own-era; built only on the real primitives) ──────────────────
def shallow_list_asof(ref_idx, history):
    """Union of the first N_SHALLOW_BLOCKS SL blocks of block_layout(ref_idx)."""
    out = set()
    for blk in block_layout(ref_idx, history, POOL)[:N_SHALLOW_BLOCKS]:
        out |= set(blk)
    return out


def spread_of(nums, ref_idx, history):
    """# distinct SL groups `nums` occupy in the since-last map as of ref_idx.
    Count is invariant to Rule-3 gap-compression relabeling."""
    slmap = since_last_map(ref_idx, history, POOL)
    return len({slmap[n] for n in nums})


def group_sizes(ref_idx, history):
    """SL-group sizes as of ref_idx (block_layout folds 'unseen' into one block)."""
    return [len(b) for b in block_layout(ref_idx, history, POOL)]


def baseline_tier_counts(sizes, pick=PICK):
    """# of pick-subsets touching exactly t distinct groups, given group sizes.
    DP: dp[t][s] = ways to pick s numbers touching t groups; take 0 or C(g,k)>=1
    per group. Validated by brute force in gate_d()."""
    dp = [[0] * (pick + 1) for _ in range(pick + 1)]
    dp[0][0] = 1
    for g in sizes:
        ndp = [row[:] for row in dp]
        for k in range(1, min(g, pick) + 1):
            ck = math.comb(g, k)
            for t in range(pick):
                for s in range(pick - k + 1):
                    if dp[t][s]:
                        ndp[t + 1][s + k] += dp[t][s] * ck
        dp = ndp
    return [dp[t][pick] for t in range(pick + 1)]


# ── GATES ─────────────────────────────────────────────────────────────────────
def gate_a(history, idx):
    di = idx.get(4687)
    got = shallow_list_asof(di, history) if di is not None else set()
    ok = got == DOC_SHALLOW_25
    print(f"[GATE A] shallow list as-of D4687 (w{di+1}) == doc 25-list: {ok}")
    return ok


def gate_b(history):
    n = len(history)
    passed = failed = 0
    exc = []
    for i in range(n):
        pred = i + 1
        if pred >= n:                     # oldest draw: no predecessor era
            continue
        if history[i]["nums"] & shallow_list_asof(pred, history):
            passed += 1
        else:
            failed += 1; exc.append(history[i]["w"])
    ok = set(exc) == KNOWN_EXCEPTIONS and (passed, failed) == (963, 6)
    print(f"[GATE B] own-era shallow anchor: pass={passed} fail={failed} "
          f"exc={sorted(exc, key=lambda w: int(w[1:]))} -> {ok}")
    return ok


def gate_c(history, idx):
    di = idx[4687]
    rows = load_r_spreads("Sheet1")
    mism = 0
    for r in rows:
        ts = r["True_Spread"]
        if ts in (None, "") or not r["nums"]:
            continue
        if spread_of(r["nums"], di, history) != int(ts):
            mism += 1
    ok = mism == 0
    print(f"[GATE C] spread mechanism (fixed-D4687) vs persisted True_Spread on "
          f"{len(rows)} R rows: mismatches={mism} -> {ok}")
    return ok


def gate_d(history, idx):
    """Brute-force the ungated Step-3 DP: exact match on synthetic eras + sum-to-C
    on the real D4687 era."""
    def brute(sizes, pick):
        gid = [g for g, sz in enumerate(sizes) for _ in range(sz)]
        c = Counter()
        for comb in combinations(range(len(gid)), pick):
            c[len({gid[i] for i in comb})] += 1
        return [c.get(t, 0) for t in range(pick + 1)]
    ok = True
    for sizes, pick in [([3,2,4,1,2],4), ([2]*6,6), ([5,1,1,3],3), ([6,4,3,2],5)]:
        ok &= baseline_tier_counts(sizes, pick) == brute(sizes, pick)
    real = baseline_tier_counts(group_sizes(idx[4687], history))
    ok &= sum(real) == C_POOL_PICK
    print(f"[GATE D] DP baseline brute-force + sum-to-C(45,6): {ok}")
    return ok


# ── STEPS 2-4 ─────────────────────────────────────────────────────────────────
def step2_spread_tiers(history):
    n = len(history)
    tiers = Counter(); low = {1: [], 2: []}
    for i in range(n):
        pred = i + 1
        if pred >= n:
            continue
        s = spread_of(history[i]["nums"], pred, history)
        tiers[s] += 1
        if s in (1, 2):
            low[s].append((history[i]["w"], sorted(history[i]["nums"])))
    ev = sum(tiers.values())
    print(f"\nSTEP 2 — own-era spread tiers, {ev} evaluable draws (oldest w{n} excluded)")
    for t in range(1, 7):
        c = tiers.get(t, 0)
        print(f"   tier {t}: {c:>4}  ({c/ev:.4%})")
    for t in (1, 2):
        print(f"   tier-{t} draws: {[w for w,_ in low[t]] or 'NONE'}"
              + "".join(f"\n      {w}: {ns}" for w, ns in low[t]))
    return tiers, ev


def step3_enrichment(history, tiers, ev):
    n = len(history)
    base = [0.0] * 7
    for i in range(n):
        pred = i + 1
        if pred >= n:
            continue
        counts = baseline_tier_counts(group_sizes(pred, history))
        for t in range(7):
            base[t] += counts[t] / C_POOL_PICK
    base = [x / ev for x in base]
    print("\nSTEP 3 — enrichment (win share / combo-space share), avg over same eras")
    print(f"   {'tier':>4} {'win_share':>10} {'combo_share':>12} {'enrichment':>11}")
    for t in range(1, 7):
        ws_ = tiers.get(t, 0) / ev; bs = base[t]
        enr = ws_ / bs if bs else float("nan")
        print(f"   {t:>4} {ws_:>10.4%} {bs:>12.4%} {enr:>11.2f}")


def step4_crosstab():
    print("\nSTEP 4 — R-row cross-tab: spread tier vs shallow-anchor status")
    def tb(sheet, label):
        c = Counter(int(r["True_Spread"]) for r in load_r_spreads(sheet)
                    if r["True_Spread"] not in (None, ""))
        tot = sum(c.values())
        t12, t3, t46 = c[1] + c[2], c[3], c[4] + c[5] + c[6]
        print(f"   {label} (n={tot}): 1-2={t12} ({t12/tot:.2%})  "
              f"3={t3} ({t3/tot:.2%})  4-6={t46} ({t46/tot:.2%})")
    tb("Shallow_Anchor_Holding", "Shallow-held")
    tb("Repeat_R", "Active Repeat")
    tb("No_Repeat_R", "Active No_Repeat")


def main():
    history = load_history()
    idx = {int(h["draw"]): i for i, h in enumerate(history) if h["draw"].isdigit()}
    print(f"Loaded {len(history)} B1 draws (w1 draw {history[0]['draw']} … w{len(history)}).")
    gates = [gate_a(history, idx), gate_b(history), gate_c(history, idx), gate_d(history, idx)]
    if not all(gates):
        print("\nHARD GATES FAILED — not tabulating."); return
    print("\nAll gates pass — tabulating.")
    tiers, ev = step2_spread_tiers(history)
    step3_enrichment(history, tiers, ev)
    step4_crosstab()


if __name__ == "__main__":
    main()
