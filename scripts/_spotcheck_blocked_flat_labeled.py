"""Spot-check for the LABELED export. Everything the original spot-check does
(three-way, adjusted for the extra header row) PLUS a boundary-label cross-check
against the actual overlap set from build_full_history (not column-index math)."""
from __future__ import annotations
import sys
from pathlib import Path
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from syndicate_core.full_history import build_full_history, draw_history_records
from syndicate_core.pipeline import b1_path, game_dirs
from syndicate_core.stacked_blocks import render_columns, column_pads, group_rail_flags
from scripts.export_blocked_flat_xlsx import cell_style

GAME, POOL, PICK = "sat", 45, 6
XLSX = ROOT / "exports" / f"blocked_flat_full_range_{GAME}_labeled.xlsx"
HDR = 4                     # labeled file has 4 header rows


def ui_truth(cell, railed):
    k = cell[0]
    if k == "num":
        n, deep = cell[1], cell[2]
        bg = ("#FFFF00" if n <= 9 else "#00B0F0" if n <= 19 else
              "#A0A0A0" if n <= 29 else "#92D050" if n <= 39 else "#FF69B4")
        if deep:
            h = bg.lstrip("#"); r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
            bg = "#%02X%02X%02X" % (int(r * .6), int(g * .6), int(b * .6)); fg = "#FFFFFF"
        else:
            fg = "#000000"
        return {"bg": bg, "fg": fg, "text": n, "railed": railed}
    if k == "hole":
        if not cell[2]:                    # inherited hole → no fill (transparent)
            return {"bg": None, "fg": None, "text": None, "railed": True} if railed else None
        return {"bg": "#FFFFFF" if cell[1] == "wall" else "#8B6F47",
                "fg": None, "text": None, "railed": railed}
    return None


def main():
    dh_df = pd.read_csv(game_dirs(GAME)["SinceLast"] / "draw_history.csv", dtype=str)
    recs, rep = build_full_history(pd.read_csv(b1_path(GAME), dtype=str), dh_df, PICK)
    n = len(recs)
    # Full export = fully-aligned view (window-global reclamation is a no-op when
    # the whole range is displayed).
    cols = render_columns(list(range(n)), recs, POOL)
    pads = column_pads(list(range(n)), recs)
    rails = [group_rail_flags(c) for c in cols]

    step, anchor, ov = rep["step"], rep["anchor_draw"], rep["overlap_n"]
    validated_ids = set(int(x) for x in rep["top_from_dh"]) | {anchor - step * i for i in range(ov)}
    cutoff = min(validated_ids)
    # INDEPENDENT ground truth: the actual draw_history explicit ids (the overlap set)
    dh_ids = {int(r["draw"]) for r in draw_history_records(dh_df) if r["draw"].lstrip("-").isdigit()}
    print(f"cutoff D{cutoff} · validated_ids==dh_ids: {validated_ids == dh_ids} · "
          f"n_validated {sum(1 for r in recs if int(r['draw']) >= cutoff)}")
    assert validated_ids == dh_ids, "report-derived validated set != draw_history ground truth"
    # every column's cutoff-label must equal dh-membership (no off-by-one anywhere)
    lbl_mm = sum(1 for r in recs if (int(r["draw"]) >= cutoff) != (int(r["draw"]) in dh_ids))
    print(f"cutoff-label vs dh-membership mismatches (all {n} cols): {lbl_mm}")
    assert lbl_mm == 0

    # ── three-way check on the D2779->D2761 slice (adjusted HDR=4) ───────────
    lo = next(i for i in range(n) if recs[i]["draw"] == "2779")
    hi = next(i for i in range(n) if recs[i]["draw"] == "2761")
    sl = list(range(lo, hi + 1))
    win = render_columns(sl, recs, POOL); win_r = [group_rail_flags(c) for c in win]
    eq = all(cols[sl[m]] == win[m] and rails[sl[m]] == win_r[m] for m in range(len(sl)))
    print(f"(a) slice == isolated window: {eq}")
    assert eq
    mm_b = 0
    for j in sl:
        for si, cell in enumerate(cols[j]):
            t, s = ui_truth(cell, rails[j][si]), cell_style(cell, rails[j][si])
            if t is None:
                mm_b += s is not None; continue
            _bg_ok = (s["bg"] is None and t["bg"] is None) or (
                s["bg"] is not None and t["bg"] is not None
                and s["bg"].upper() == t["bg"].upper())
            if not (s and _bg_ok and s["font"] == t["fg"]
                    and s["text"] == t["text"] and s["railed"] == t["railed"]):
                mm_b += 1
    print(f"(b) cell_style vs UI-truth mismatches: {mm_b}")
    assert mm_b == 0

    expected = {}
    for j in sl:
        for si, cell in enumerate(cols[j]):
            t = ui_truth(cell, rails[j][si])
            if t is None:
                continue
            exrow = HDR + pads[j] + si + 1
            argb = ("FF" + t["bg"].lstrip("#").upper()) if t["bg"] is not None else None
            expected[(exrow, j + 1)] = (argb, t["text"], t["railed"])
    wb = openpyxl.load_workbook(XLSX, read_only=True); ws = wb.active
    min_c, max_c = sl[0] + 1, sl[-1] + 1
    checked = fill_mm = val_mm = rail_mm = 0; seen = set(); exrow = HDR
    for row in ws.iter_rows(min_row=HDR + 1, min_col=min_c, max_col=max_c):
        exrow += 1
        for ci, c in enumerate(row):
            key = (exrow, min_c + ci)
            if key not in expected:
                continue
            argb, val, railed = expected[key]
            got = c.fill.fgColor.rgb if (c.fill and c.fill.patternType) else None
            fill_mm += got != argb
            val_mm += (c.value if c.value is not None else None) != val
            hr = bool(c.border and c.border.left and c.border.left.style
                      and c.border.left.color and str(c.border.left.color.rgb).endswith("777777"))
            rail_mm += hr != railed
            checked += 1; seen.add(key)
    print(f"(c) checked {checked} · fill_mm {fill_mm} · val_mm {val_mm} · rail_mm {rail_mm} · "
          f"missing {len(expected) - len(seen)}")
    assert fill_mm == 0 and val_mm == 0 and rail_mm == 0 and len(seen) == len(expected)

    # ── NEW: 4th-row label + orange strip, boundary columns (read back from file) ──
    # read the whole 4th header row (excel row 4) once
    label_row, orange_top = {}, {}
    for row in ws.iter_rows(min_row=4, max_row=4):
        for ci, c in enumerate(row):
            label_row[ci + 1] = c.value
            top = c.border.top if c.border else None
            orange_top[ci + 1] = bool(top and top.style and top.color
                                      and str(top.color.rgb).endswith("FFA500"))
    wb.close()
    col_by_draw = {int(recs[j]["draw"]): j + 1 for j in range(n)}
    print("    boundary columns (label read from file · orange strip · dh-truth):")
    ok = True
    for d in (4407, 4405, 4403, 4401, 4399):
        excol = col_by_draw[d]
        want_val = d in dh_ids                       # ground truth, not index math
        got_label = label_row.get(excol)
        want_label = "validated" if want_val else "extrapolated"
        want_orange = not want_val                   # orange strip on extrapolated only
        got_orange = orange_top.get(excol, False)
        good = (got_label == want_label) and (got_orange == want_orange)
        ok = ok and good
        print(f"      D{d}: file='{got_label}' orange={got_orange} | "
              f"dh_truth={want_label} want_orange={want_orange} -> {'OK' if good else 'FAIL'}")
    assert ok, "boundary label/strip mismatch"
    print("LABELED SPOT-CHECK PASS: colours cell-for-cell + labels/strip correct at the boundary.")


if __name__ == "__main__":
    main()
